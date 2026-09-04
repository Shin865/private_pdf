"""
PowerPoint conversion service for PDF Tool.
Handles PDF → PowerPoint (.pptx) conversion with editable and exact modes.
"""
import io
import logging
from typing import Optional

import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from app.config import RENDER_DPI_STANDARD, RENDER_DPI_HIGH

logger = logging.getLogger("pdf-tool.powerpoint-service")


def _get_slide_dimensions(pdf_page_width: float, pdf_page_height: float):
    """
    Convert PDF page dimensions (in points) to PowerPoint slide dimensions.
    Returns (width_emu, height_emu).
    """
    # PDF points to EMU (1 point = 12700 EMU)
    width_emu = int(pdf_page_width * 12700)
    height_emu = int(pdf_page_height * 12700)
    return width_emu, height_emu


def _add_editable_slide(prs: Presentation, doc: fitz.Document, page_idx: int):
    """
    Add a slide with editable text boxes and images extracted from PDF page.
    """
    page = doc[page_idx]
    rect = page.rect

    # Add blank slide
    blank_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(blank_layout)

    # Scale factor from PDF points to slide EMUs
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    scale_x = slide_width / (rect.width * 12700)
    scale_y = slide_height / (rect.height * 12700)
    scale = min(scale_x, scale_y)

    # Extract text blocks
    blocks = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)["blocks"]

    for block in blocks:
        if block["type"] == 0:  # Text block
            x0, y0, x1, y1 = block["bbox"]

            # Convert to EMU with scaling
            left = int(x0 * 12700 * scale)
            top = int(y0 * 12700 * scale)
            width = int((x1 - x0) * 12700 * scale)
            height = int((y1 - y0) * 12700 * scale)

            # Ensure minimum dimensions
            width = max(width, Inches(0.5))
            height = max(height, Inches(0.2))

            # Clamp to slide bounds
            if left + width > slide_width:
                width = slide_width - left
            if top + height > slide_height:
                height = slide_height - top
            if left < 0:
                left = 0
            if top < 0:
                top = 0

            try:
                txBox = slide.shapes.add_textbox(left, top, width, height)
                tf = txBox.text_frame
                tf.word_wrap = True

                first_para = True
                for line in block.get("lines", []):
                    line_text = ""
                    line_font_size = 12
                    line_is_bold = False
                    line_color = (0, 0, 0)

                    for span in line.get("spans", []):
                        line_text += span["text"]
                        line_font_size = span.get("size", 12)
                        flags = span.get("flags", 0)
                        line_is_bold = bool(flags & 2**4)
                        color_int = span.get("color", 0)
                        r = (color_int >> 16) & 0xFF
                        g = (color_int >> 8) & 0xFF
                        b = color_int & 0xFF
                        line_color = (r, g, b)

                    if line_text.strip():
                        if first_para:
                            para = tf.paragraphs[0]
                            first_para = False
                        else:
                            para = tf.add_paragraph()

                        run = para.add_run()
                        run.text = line_text.strip()
                        run.font.size = Pt(min(max(line_font_size * scale, 4), 72))
                        run.font.bold = line_is_bold
                        run.font.color.rgb = RGBColor(*line_color)
            except Exception as e:
                logger.warning(f"Could not add text box: {e}")

    # Extract and add images
    for img_idx, img in enumerate(page.get_images(full=True)):
        xref = img[0]
        try:
            base_image = doc.extract_image(xref)
            if base_image:
                img_bytes = base_image["image"]
                img_stream = io.BytesIO(img_bytes)

                # Try to get image position from page
                img_rects = page.get_image_rects(xref)
                if img_rects:
                    img_rect = img_rects[0]
                    left = int(img_rect.x0 * 12700 * scale)
                    top = int(img_rect.y0 * 12700 * scale)
                    width = int((img_rect.x1 - img_rect.x0) * 12700 * scale)
                    height = int((img_rect.y1 - img_rect.y0) * 12700 * scale)
                else:
                    left = Inches(1)
                    top = Inches(1)
                    width = Inches(4)
                    height = Inches(3)

                slide.shapes.add_picture(img_stream, left, top, width, height)
        except Exception as e:
            logger.warning(f"Could not add image to slide: {e}")


def _add_exact_slide(prs: Presentation, doc: fitz.Document, page_idx: int, dpi: int):
    """
    Add a slide with the PDF page rendered as a full-page image (exact layout).
    """
    page = doc[page_idx]

    # Render page as image
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat)
    img_bytes = pix.tobytes("png")

    # Add blank slide
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Add image to cover entire slide
    img_stream = io.BytesIO(img_bytes)
    slide.shapes.add_picture(
        img_stream, 0, 0, prs.slide_width, prs.slide_height
    )


def pdf_to_powerpoint(
    pdf_bytes: bytes,
    mode: str = "exact",
    pages: list[int] = None,
    resolution: str = "standard",
) -> bytes:
    """
    Convert PDF to PowerPoint (.pptx).

    Args:
        pdf_bytes: PDF file content
        mode: "editable" or "exact"
        pages: optional list of page numbers (1-indexed). None = all pages
        resolution: "standard" (150 DPI) or "high" (300 DPI) — for exact mode

    Returns:
        PPTX file content as bytes
    """
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    total = len(doc)

    target_pages = pages if pages else list(range(1, total + 1))
    target_pages = [p for p in target_pages if 1 <= p <= total]

    dpi = RENDER_DPI_HIGH if resolution == "high" else RENDER_DPI_STANDARD

    # Create presentation
    prs = Presentation()

    # Set slide dimensions based on first target page
    if target_pages:
        first_page = doc[target_pages[0] - 1]
        rect = first_page.rect
        prs.slide_width, prs.slide_height = _get_slide_dimensions(rect.width, rect.height)

    for p in target_pages:
        page_idx = p - 1

        if mode == "editable":
            _add_editable_slide(prs, doc, page_idx)
        else:
            _add_exact_slide(prs, doc, page_idx, dpi)

        logger.info(f"Added slide for page {p} (mode={mode})")

    doc.close()

    # Save to bytes
    output = io.BytesIO()
    prs.save(output)
    result = output.getvalue()

    logger.info(
        f"PDF → PowerPoint: {len(target_pages)} slides, mode={mode}, "
        f"resolution={resolution}, {len(result)} bytes"
    )
    return result
